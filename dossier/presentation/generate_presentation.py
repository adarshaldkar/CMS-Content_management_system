"""
The Gazette CMS — Presentation Generator Script
Generates a 15-slide widescreen (.pptx) engineering presentation deck using python-pptx.
"""

import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Theme Palette
    COLOR_BG = RGBColor(15, 23, 42)       # Slate 900
    COLOR_CARD = RGBColor(30, 41, 59)     # Slate 800
    COLOR_CYAN = RGBColor(56, 189, 248)   # Cyan / Sky 400
    COLOR_CORAL = RGBColor(251, 146, 60)  # Orange / Coral 400
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_SLATE = RGBColor(148, 163, 184) # Slate 400
    COLOR_GREEN = RGBColor(52, 211, 153)  # Emerald 400

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="THE GAZETTE CMS — ENGINEERING DEFENSE"):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_CYAN

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_WHITE
        p1.space_before = Pt(4)

    def add_card(slide, left, top, width, height, title, body_bullets, accent_color=COLOR_CYAN):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = accent_color
        card.line.width = Pt(1.5)

        tb = slide.shapes.add_textbox(Inches(left + 0.3), Inches(top + 0.25), Inches(width - 0.6), Inches(height - 0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = accent_color

        for item in body_bullets:
            p = tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size = Pt(12)
            p.font.color.rgb = COLOR_SLATE
            p.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)
    
    # Hero Title Box
    tb = slide1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(3.8))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "THE GAZETTE"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    
    p2 = tf.add_paragraph()
    p2.text = "Enterprise Polyglot Publication Operating System & Editorial Studio"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_WHITE
    p2.space_before = Pt(10)

    p3 = tf.add_paragraph()
    p3.text = "Decoupled React 19 Frontend | NestJS Enterprise Core | PostgreSQL + MongoDB + Redis Mesh | Zero-Trust OWASP Hardening"
    p3.font.size = Pt(14)
    p3.font.color.rgb = COLOR_SLATE
    p3.space_before = Pt(14)

    # Footer Metadata
    tb_foot = slide1.shapes.add_textbox(Inches(1.2), Inches(5.6), Inches(10.9), Inches(1.2))
    tf_foot = tb_foot.text_frame
    p_meta = tf_foot.paragraphs[0]
    p_meta.text = "Candidate / Presenter: Adarsh Patel  |  Repository: adarshaldkar/CMS-Content_management_system  |  October 2026"
    p_meta.font.size = Pt(12)
    p_meta.font.bold = True
    p_meta.font.color.rgb = COLOR_CORAL

    # -------------------------------------------------------------
    # SLIDE 2: Problem Space & Industry Gap
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "The Problem Space & Modern Editorial Bottlenecks")

    add_card(slide2, 0.8, 1.8, 3.7, 4.8, "1. Monolithic Storage Fatigue", [
        "Traditional CMS (WordPress, Drupal) cram relational metadata and complex JSON articles into a single SQL table.",
        "High-velocity writing sessions cause database table locking, cache invalidation storms, and query bottlenecks.",
        "Lack of native document modeling for rich nested rich-text block hierarchies."
    ], COLOR_CORAL)

    add_card(slide2, 4.8, 1.8, 3.7, 4.8, "2. Editorial Fragility & Content Loss", [
        "Authors face browser crashes, network drops, or expired JWT sessions during long writing workflows.",
        "Lack of local client-side resilience forces constant re-drafting.",
        "Absence of granular multi-stage review (Authors vs. Editors vs. Admins) leads to accidental unverified publishing."
    ], COLOR_CORAL)

    add_card(slide2, 8.8, 1.8, 3.7, 4.8, "3. Insecure Ingestion & Ingestion Lag", [
        "External embed scrapers and media upload endpoints frequently suffer from Server-Side Request Forgery (SSRF).",
        "Synchronous media processing blocks the main API thread, ballooning HTTP response times.",
        "Absence of strict brute-force rate-limiting invites credential stuffing."
    ], COLOR_CORAL)

    # -------------------------------------------------------------
    # SLIDE 3: Product Vision & Engineering Pillars
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "The Gazette Solution & Architectural Pillars")

    add_card(slide3, 0.8, 1.8, 5.7, 2.3, "Polyglot Persistence Mesh", [
        "PostgreSQL 16 for strict ACID user credentials, roles, and audit trails.",
        "MongoDB 7 for flexible, polymorphic ProseMirror content trees and revisions.",
        "Redis 7 for microsecond sliding-window throttlers and session revocations."
    ], COLOR_CYAN)

    add_card(slide3, 6.8, 1.8, 5.7, 2.3, "Zero-Data-Loss Writing Studio", [
        "TipTap / ProseMirror rich editor with real-time slash commands and floating ribbons.",
        "Dual-tier persistence: Keystroke-level IndexedDB local buffer + debounced server sync.",
        "Side-by-side revision conflict resolver upon re-entry."
    ], COLOR_CYAN)

    add_card(slide3, 0.8, 4.4, 5.7, 2.4, "Asynchronous Background Processing", [
        "BullMQ distributed message queues backed by Redis.",
        "Offloaded video transcoding, OpenGraph scraping, and automated publishing slots.",
        "Prevents API gateway thread starvation during heavy media operations."
    ], COLOR_GREEN)

    add_card(slide3, 6.8, 4.4, 5.7, 2.4, "Enterprise Defense-in-Depth", [
        "Zero-Trust SSRF validator blocking private subnets and AWS metadata APIs.",
        "Dual-Token JWT lifecycle (15-min access + 7-day family-rotated refresh token).",
        "Verified through live Burp Suite Pro Collaborator OOB penetration tests."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 4: As-Built Product Requirements Document (PRD)
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "As-Built Product Requirements & Acceptance Criteria")

    add_card(slide4, 0.8, 1.8, 5.7, 4.8, "Core Implemented Requirements", [
        "FR-01 (Auth): Argon2id password encryption with dual JWT cookie lifecycle.",
        "FR-02 (RBAC): Strict 5-tier role hierarchy (SUPER_ADMIN to READER).",
        "FR-03 (Studio): WYSIWYG TipTap editing with dynamic slash (/) palette.",
        "FR-04 (Recovery): Local IndexedDB autosave buffer with collision modals.",
        "FR-05 (Media): Cloudinary CDN upload + zero-trust SSRF embed scraper.",
        "FR-06 (Workflows): 5-state editorial pipeline (Draft to Published)."
    ], COLOR_CYAN)

    add_card(slide4, 6.8, 1.8, 5.7, 4.8, "Non-Functional & Operational Standards", [
        "NFR-01 (Latency): Sub-100ms cached reading responses; <50ms p95 API latency.",
        "NFR-02 (Security): OWASP Top 10 compliance (Zero SQLi, SSRF, XSS, CSRF).",
        "NFR-03 (Resilience): Graceful fallback to local storage if Redis or secondary DB settles.",
        "NFR-04 (Footprint): Docker developer mesh constrained to <=450 MB RAM budget.",
        "Acceptance Proof: All journeys validated via automated Playwright E2E suites."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 5: UI/UX & Editorial Studio Experience
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "UI/UX Engineering & Editorial Studio Ergonomics")

    add_card(slide5, 0.8, 1.8, 3.7, 4.8, "Distraction-Free Canvas", [
        "TipTap ProseMirror engine tailored for long-form journalistic writing.",
        "Context-sensitive floating bubble toolbar on text selection.",
        "Slash Command Palette (/) for keyboard-only insertion of tables, code, and callouts."
    ], COLOR_CYAN)

    add_card(slide5, 4.8, 1.8, 3.7, 4.8, "Active Document Navigation", [
        "Real-time Document Outline Drawer generated dynamically from H1-H3 DOM tags.",
        "Pre-Publish Checklist validating excerpt, slug uniqueness, and cover assets.",
        "Side-by-side visual diff comparison for editors reviewing revisions."
    ], COLOR_CYAN)

    add_card(slide5, 8.8, 1.8, 3.7, 4.8, "State & Collision Handling", [
        "Debounced 500ms auto-commit to browser IndexedDB.",
        "Unsaved changes modal intercepts accidental navigation.",
        "Visual recovery dialog comparing local timestamp with remote server snapshot."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 6: System Architecture (C4 Topology)
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "Full-Stack System Architecture & C4 Topology")

    add_card(slide6, 0.8, 1.8, 3.7, 4.8, "Client Tier (Vite / React 19)", [
        "Single-Page Application with React Router 7 and TanStack Query.",
        "State management with Zustand and reactive UI motion via Motion.",
        "Unified API Service Gateway communicating over secure HTTPS/WSS.",
        "Driver.js guided onboarding for newly invited authors."
    ], COLOR_CYAN)

    add_card(slide6, 4.8, 1.8, 3.7, 4.8, "API Gateway (NestJS Core)", [
        "Modular domain controllers with class-validator DTO pipelines.",
        "Global Guards: JwtAuthGuard, RolesGuard, CsrfGuard, ThrottlerGuard.",
        "Centralized Global Exception Filter with unique tracing reference IDs.",
        "Swagger / OpenAPI automated documentation contracts."
    ], COLOR_CYAN)

    add_card(slide6, 8.8, 1.8, 3.7, 4.8, "Persistence & Processing Mesh", [
        "PostgreSQL: ACID relational model for users, permissions, and audit logs.",
        "MongoDB: High-velocity document store for stories and block trees.",
        "Redis: In-memory sliding-window throttlers and BullMQ queues.",
        "ClickHouse: Real-time telemetry columnar database."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 7: Database Design — Polyglot Persistence
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "Database Design: Polyglot Persistence Mesh")

    add_card(slide7, 0.8, 1.8, 5.7, 2.3, "PostgreSQL 16 (Relational Core)", [
        "users: UUID PK, unique email, Argon2id hash, role enum, isActive flag.",
        "refresh_sessions: Token family ID, cryptographic JTI, revocation state.",
        "categories & tags: Relational taxonomy and editorial taxonomy trees."
    ], COLOR_CYAN)

    add_card(slide7, 6.8, 1.8, 5.7, 2.3, "MongoDB 7 (Document Store)", [
        "articles: polymorphic JSON block trees, nested revisions, autosave buffer.",
        "comments: hierarchical threaded discussion trees.",
        "publishing_slots: reserved scheduling records and author lock state."
    ], COLOR_CYAN)

    add_card(slide7, 0.8, 4.4, 5.7, 2.4, "Redis 7 (Ephemeral Cache & Queues)", [
        "throttle:default & throttle:auth-sensitive sliding window counters.",
        "session:revoked:<jti> blacklist sets for instant logout propagation.",
        "bull:editorial-queue:* for persistent background job payloads."
    ], COLOR_CORAL)

    add_card(slide7, 6.8, 4.4, 5.7, 2.4, "ClickHouse (Telemetry Engine)", [
        "Real-time event logging: article_views, reading_duration, scroll_depth.",
        "Columnar compression enables querying 10M+ rows in sub-second time.",
        "Powers author reading analytics dashboard without SQL database lag."
    ], COLOR_CORAL)

    # -------------------------------------------------------------
    # SLIDE 8: Backend Engineering & API Design
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Backend Engineering & API Design Standards")

    add_card(slide8, 0.8, 1.8, 3.7, 4.8, "Domain-Driven Modularity", [
        "AuthModule: Authentication, token refresh, Turnstile captcha.",
        "ArticlesModule: Editorial workflows, revision forks, search.",
        "EmbedsModule: OpenGraph scrapers, YouTube / Vimeo providers.",
        "JobsModule: BullMQ queue workers and scheduler cron."
    ], COLOR_CYAN)

    add_card(slide8, 4.8, 1.8, 3.7, 4.8, "Guard & Pipeline Flow", [
        "Incoming Request -> CORS / Helmet security headers.",
        "CsrfGuard: Double-Submit Cookie verification for state-changing calls.",
        "ThrottlerGuard: Multi-tier sliding window rate checks.",
        "JwtAuthGuard: HMAC verification and active user status lookup."
    ], COLOR_CYAN)

    add_card(slide8, 8.8, 1.8, 3.7, 4.8, "Enterprise Resilience", [
        "Centralized Global Exception Filter maps all errors to structured RFC-7807 JSON.",
        "Generates traceable reference IDs (ERR-YYYYMMDD-XXXXXX) for debugging.",
        "Strict DTO validation with whitelist and forbidNonWhitelisted flags."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 9: Asynchronous Pipelines (BullMQ & Jobs)
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9)
    add_header(slide9, "Asynchronous Pipelines & Scheduled Publishing")

    add_card(slide9, 0.8, 1.8, 5.7, 4.8, "Publishing Slot Architecture", [
        "Problem: Authors need to release stories at peak reader engagement hours without manual intervention.",
        "Solution: BullMQ repeatable delayed jobs scheduled on Redis.",
        "Step 1: Editor reserves slot -> Record stored in MongoDB.",
        "Step 2: Delayed BullMQ job enqueued with exact future timestamp.",
        "Step 3: Worker awakens at target timestamp -> Transitions article to PUBLISHED and busts cache.",
        "Resilience: Automatic retries with exponential backoff on transient database failures."
    ], COLOR_CYAN)

    add_card(slide9, 6.8, 1.8, 5.7, 4.8, "Media & Document Processing Pipeline", [
        "Heavy file parsing (DOCX, PDF, image transcoding) is kept strictly off the HTTP thread.",
        "Incoming upload generates signed Cloudinary URL or streams to local disk fallback.",
        "Document Parser extracts structure and transforms headings into TipTap JSON nodes.",
        "Media queue generates responsive thumbnails and WebP optimized variants asynchronously."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 10: Security Architecture & Threat Model
    # -------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10)
    add_header(slide10, "Security Architecture & OWASP Defense Matrix")

    add_card(slide10, 0.8, 1.8, 3.7, 4.8, "Authentication & Tokens", [
        "Argon2id password hashing (memory cost: 64MB, 3 iterations).",
        "Short-lived 15-min JWT access token in HTTP-only, SameSite cookies.",
        "7-day refresh token with Family Token Rotation (detects token reuse and instantly revokes family)."
    ], COLOR_CORAL)

    add_card(slide10, 4.8, 1.8, 3.7, 4.8, "SSRF Zero-Trust Guard", [
        "Custom SSRF Validator (ssrf-validator.util.ts) executes prior to outbound HTTP fetches.",
        "Resolves DNS hostnames and verifies target IP against zero-trust blocklists.",
        "Blocks RFC 1918 private subnets (10.*, 172.16.*, 192.168.*), loopback (127.0.0.1), and AWS metadata (169.254.169.254)."
    ], COLOR_CORAL)

    add_card(slide10, 8.8, 1.8, 3.7, 4.8, "Application Hardening", [
        "CsrfGuard: Enforces double-submit cookie verification on all POST/PUT/DELETE requests.",
        "DOMPurify sanitization eliminates Stored and Reflected XSS vectors.",
        "Redis sliding-window rate limiters block automated credential brute-forcing."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 11: Penetration Testing & Burp Suite Audit
    # -------------------------------------------------------------
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11)
    add_header(slide11, "Live Penetration Testing & Burp Suite Verification")

    add_card(slide11, 0.8, 1.8, 3.7, 4.8, "1. Out-of-Band (OAST) Blind SSRF", [
        "Tested POST /api/embeds/scrape with Burp Collaborator payload (oastify.com).",
        "Result: Collaborator captured incoming DNS lookup and HTTP verification.",
        "Validation: Attempted loopback (127.0.0.1) and AWS metadata -> blocked with HTTP 400 Bad Request."
    ], COLOR_CYAN)

    add_card(slide11, 4.8, 1.8, 3.7, 4.8, "2. Rate-Limiting Verification", [
        "Targeted POST /api/auth/login using Burp Repeater & Intruder rapid bursts.",
        "Result: Attempts 1-5 succeeded. Attempt 6 triggered HTTP 429 Too Many Requests.",
        "Header Verified: retry-after-auth-sensitive: 56; verified Redis sliding-window enforcement."
    ], COLOR_CYAN)

    add_card(slide11, 8.8, 1.8, 3.7, 4.8, "3. JWT Signature Tampering", [
        "Used JWT Editor extension to tamper payload claims ('role': 'AUTHOR' -> 'ADMIN').",
        "Forwarded without valid HMAC secret signature.",
        "Result: Server rejected request with HTTP 401 Unauthorized: Session expired or invalid token."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 12: Testing Strategy & Verification Metrics
    # -------------------------------------------------------------
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12)
    add_header(slide12, "Testing Strategy & Automated Quality Assurance")

    add_card(slide12, 0.8, 1.8, 3.7, 4.8, "Unit Testing (Vitest)", [
        "High-coverage test suites for isolated utility algorithms.",
        "Cryptographic token signing, expiration logic, and cookie parsers.",
        "SSRF IP blocklist evaluation across all IPv4/IPv6 private ranges.",
        "Fast feedback loop during development (<2.5s execution)."
    ], COLOR_CYAN)

    add_card(slide12, 4.8, 1.8, 3.7, 4.8, "Integration Tests", [
        "Validates NestJS dependency injection and database repository bindings.",
        "Tests transactional user creation and MongoDB article persistence.",
        "Verifies BullMQ job lifecycle: enqueue, delay, and worker settlement."
    ], COLOR_CYAN)

    add_card(slide12, 8.8, 1.8, 3.7, 4.8, "E2E Golden Journey (Playwright)", [
        "Automated multi-role browser journeys in headless Chromium.",
        "Journey 1: SuperAdmin user onboarding and taxonomy seeding.",
        "Journey 2: Author drafting, image embed, and autosave verification.",
        "Journey 3: Editor review, approval, and live article publication."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 13: Operational Demonstration & Live Workflows
    # -------------------------------------------------------------
    slide13 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide13)
    add_header(slide13, "Operational Demonstration & Verified Workflows")

    add_card(slide13, 0.8, 1.8, 5.7, 4.8, "Live Operational Capabilities", [
        "Editorial Creation: Authors compose with slash commands and floating toolbars in sub-second response times.",
        "Autosave Resilience: Verified recovery of 100% of unsaved content via IndexedDB after simulated tab crash.",
        "Editorial Governance: Editors review revision diffs side-by-side prior to approval.",
        "Reader Engagement: Clean, responsive reading canvas with live comments, likes, and reading progress bars.",
        "Docker Mesh: Active containers (MongoDB, Postgres, Redis) operating within ~335 MB RAM budget."
    ], COLOR_CYAN)

    add_card(slide13, 6.8, 1.8, 5.7, 4.8, "Key Verification Proof Points", [
        "Zero-Trust Ingestion: Server successfully blocks internal network pivots.",
        "Brute-Force Resistance: Redis throttler defends authentication against credential stuffing.",
        "Tamper-Proof Authorization: Role escalation attempts fail at server-side guards.",
        "Multi-Engine Cohesion: PostgreSQL credentials and MongoDB article trees remain synchronized via transactional hooks."
    ], COLOR_GREEN)

    # -------------------------------------------------------------
    # SLIDE 14: Technical Trade-offs & Future Scope
    # -------------------------------------------------------------
    slide14 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide14)
    add_header(slide14, "Technical Trade-offs & Future Engineering Roadmap")

    add_card(slide14, 0.8, 1.8, 5.7, 4.8, "Current Engineering Trade-offs", [
        "Dual-Database Consistency: Maintaining referential integrity between PostgreSQL users and MongoDB stories requires application-level hooks rather than database foreign keys.",
        "Single-Tenant Architecture: Designed for single media enterprise; currently requires independent deployments for multiple publications.",
        "Memory Budget Constraints: Secondary analytics engines (ClickHouse) are maintained under an optional Docker profile to prevent workstation RAM saturation."
    ], COLOR_CORAL)

    add_card(slide14, 6.8, 1.8, 5.7, 4.8, "Future Engineering Roadmap", [
        "Multi-Tenant SaaS Partitioning: Organization-level schema isolation and vanity domain routing.",
        "Real-Time Collaborative Editing: Integration of Yjs CRDTs over WebSockets for multi-author live typing (Google Docs style).",
        "AI Editorial Co-Pilot: Integrated LLM assistants for SEO metadata generation, headline suggestions, and style auditing.",
        "Edge Caching: Cloudflare Workers integration for global sub-20ms reader delivery."
    ], COLOR_CYAN)

    # -------------------------------------------------------------
    # SLIDE 15: Conclusion & Q&A
    # -------------------------------------------------------------
    slide15 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide15)
    
    tb = slide15.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(3.8))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "CONCLUSION & ENGINEERING SUMMARY"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    
    p2 = tf.add_paragraph()
    p2.text = "The Gazette proves that enterprise publishing demands far more than basic CRUD. By decoupling editorial ergonomics from polyglot persistence and applying zero-trust defensive security, the platform achieves both authoring excellence and operational resilience."
    p2.font.size = Pt(16)
    p2.font.color.rgb = COLOR_WHITE
    p2.space_before = Pt(14)

    p3 = tf.add_paragraph()
    p3.text = "Key Contributions: Zero-loss IndexedDB recovery | Polyglot (Postgres+Mongo+Redis) Mesh | BullMQ Async Queues | Live Burp Suite Penetration Verified"
    p3.font.size = Pt(14)
    p3.font.color.rgb = COLOR_GREEN
    p3.space_before = Pt(14)

    tb_qa = slide15.shapes.add_textbox(Inches(1.2), Inches(5.6), Inches(10.9), Inches(1.2))
    tf_qa = tb_qa.text_frame
    p_qa = tf_qa.paragraphs[0]
    p_qa.text = "Thank you! Open for Questions & Architecture Review."
    p_qa.font.size = Pt(20)
    p_qa.font.bold = True
    p_qa.font.color.rgb = COLOR_CORAL

    output_path = os.path.join(os.path.dirname(__file__), "The_Gazette_CMS_Final_Presentation.pptx")
    prs.save(output_path)
    print(f"[SUCCESS] Presentation generated successfully at: {output_path}")

if __name__ == "__main__":
    create_presentation()
